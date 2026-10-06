# brain用に作成、グラフ生成キット
import pandas as pd
from datetime import datetime
import openpyxl as px
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.axis import DateAxis

#file名を入力させ、保存
count=0
time_row_start=[]
time_row_end=[]
title_list=["平均累積報酬","エピソード毎のステップ数","オブジェクト接触回数","転倒回数","カウント超過回数","成功回数","Agentの数"]
y_axis_list=["reward","step","touch","roll","count_over","success","Agent"]
y_axis_min=[-20,0,0,0,0,0,0]
y_axis_max=[100,11000,150,25,25,25,25]
y_axis_interval=[20,2000,20,5,5,5,5]
sell_width=["J1","J16","S1","S16","AB1","AB16","S31"]

filepath=input("file名は何ですか→")

excel_file=px.load_workbook(filepath)
# excel_sheet=excel_file['time1']
excel_sheet=excel_file.active

for i in range(1,100000):
    if excel_sheet.cell(row=i,column=1).value=="episode":
        time_row_start.append(i)
        count=0
    if excel_sheet.cell(row=i,column=1).value==None:
        count=count+1
        time_row_end.append(i)
        break

for i in range(0,len(time_row_start)):
    for k in range(0,len(title_list)):
        # chart = px.chart.LineChart()
        chart = ScatterChart()
        # y軸のデータ
        data=px.chart.Reference(excel_sheet,k+2,time_row_start[i],k+2,time_row_end[i]-1)#左上列、行、右下列行
        # x軸のデータ
        categories=px.chart.Reference(excel_sheet,1,time_row_start[i]+1,1,time_row_end[i]-1)#左上列、行、右下列行
        series=Series(data,categories,title_from_data=True)
        chart.series.append(series)
        # chart.add_data(data,titles_from_data=True)
        # chart.set_categories(categories)
        chart.style = 3
        num=str(i+1)
        chart.title = title_list[k]
        chart.y_axis.title = y_axis_list[k]
        chart.x_axis.title = "エピソード"
        chart.height = 9
        chart.width = 16
        chart.x_axis.scaling.min = 0
        chart.x_axis.scaling.max = time_row_end[0]
        chart.y_axis.scaling.min = y_axis_min[k]
        chart.y_axis.scaling.max = y_axis_max[k]
        chart.x_axis.majorUnit = 100
        chart.y_axis.majorUnit = y_axis_interval[k]

        excel_sheet.add_chart(chart,""+sell_width[k]+"")

excel_file.save("gragh_"+filepath)